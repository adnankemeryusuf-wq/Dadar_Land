import os
import requests
import qrcode
import openpyxl
from datetime import datetime
from openpyxl.styles import Font, Alignment, PatternFill

# --- QINDAA'INA (CONFIG) ---
USER_NAME = "admin"
PASS_WORD = "1234"
BOT_TOKEN = "8586015354:AAEliISf-RtoeJ8anbVLapY3NBm7hz8dZWI"
CHAT_ID_MANAGER = "568248052"
DATA_FILE = "dadar_final_report.txt"

# --- FUNKSHIINIIWWAN ERGAA ---
def send_telegram(file_path, file_type="doc", caption=""):
    url = f"https://api.telegram.org/bot{BOT_TOKEN}/"
    url += "sendDocument" if file_type == "doc" else "sendPhoto"
    try:
        with open(file_path, 'rb') as f:
            payload = {'chat_id': CHAT_ID_MANAGER, 'caption': caption}
            files = {'document' if file_type == "doc" else 'photo': f}
            requests.post(url, data=payload, files=files, timeout=20)
    except: 
        print(f"[!] Ergaan Telegram hin darbine. Internet kee itti fufi.")

# --- GALMEESSI (NEW ENTRY) ---
def galmeessi():
    print("\n" + "="*20 + " GALMEE HAARAA " + "="*20)
    ad = input("Maqaa AD: "); bad = input("Bilbila AD: ")
    gs = input("Gosa Tajaajilaa: "); og = input("Maqaa Ogeessa: ")
    bog = input("Bilbila Ogeessa: "); gb = input("Beellama (YYYY-MM-DD): "); sb = input("Sa'aatii: ")
    
    k_vals = []
    for kt in ["Kartaa", "User", "Jij_Maqaa", "Lizio", "TOT", "Gibira"]:
        while True:
            v = input(f"{kt}: ") or "0"
            try: k_vals.append(float(v)); break
            except: print("[!] Maaloo lakkoofsa galchi!")
            
    waligala = sum(k_vals)
    now_str = datetime.now().strftime("%Y-%m-%d %H:%M")
    
    # Save & QR
    qr_file = f"QR_{ad.replace(' ', '_')}.png"
    qrcode.make(f"AD: {ad}\nKafaltii: {waligala}\nBeellama: {gb}").save(qr_file)
    send_telegram(qr_file, "photo", f"✅ QR Haaraa: {ad}\n💰 Waligala: {waligala} ETB")
    
    # Line format: 16 columns (indices up to 15) to match your Excel logic
    k_str = "|".join(map(str, k_vals))
    line = f"{now_str}|{ad}|{bad}|-|-|{gs}|{og}|{bog}|{gb} {sb}|{k_str}|{waligala}\n"
    with open(DATA_FILE, "a", encoding="utf-8") as f: f.write(line)
    print(f"\n[✓] Galmeeffameera! Waligala: {waligala} ETB")

# --- FUNKSHIINII GABAASAA (EXCEL) ---
def uumi_excel(days, label):
    if not os.path.exists(DATA_FILE):
        print("[!] Ragaan kuufame hin jiru."); return
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = label
    headers = ["Guyyaa", "Abbaa Dhimmaa", "Gosa Tajaajilaa", "Ogeessa", "Beellama", "Waligala"]
    ws.append(headers)
    
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    
    now = datetime.now()
    count = 0
    total_revenue = 0
    
    with open(DATA_FILE, "r", encoding="utf-8") as f:
        for line in f:
            p = line.strip().split("|")
            try:
                dt = datetime.strptime(p[0], "%Y-%m-%d %H:%M")
                if (now - dt).days <= days:
                    # Index 15 refers to the last element (Waligala)
                    ws.append([p[0], p[1], p[5], p[6], p[8], p[-1]])
                    total_revenue += float(p[-1])
                    count += 1
            except: continue
            
    if count == 0:
        print(f"[!] Gabaasa {label} irratti ragaan argame hin jiru."); return

    file_name = f"Gabaasa_{label}.xlsx"
    wb.save(file_name)
    send_telegram(file_name, "doc", f"📊 Gabaasa {label}\n💰 Waligala Galii: {total_revenue} ETB")
    print(f"[✓] Gabaasni {label} hoggantatti ergameera.")

# --- MAIN MENU ---
if __name__ == "__main__":
    print("\n" + "*"*35 + "\n W/BULCHINSA LAFAA MAGAALAA DADAR\n" + "*"*35)
    u, p = input("User: "), input("Pass: ")
    if u == USER_NAME and p == PASS_WORD:
        while True:
            print("\n1. Galmee Haaraa\n2. Gabaasa Excel (Telegram)\n3. Barbaadi (Search)\n4. Exit")
            c = input("\nFilannoo kee: ")
            if c == '1': galmeessi()
            elif c == '2':
                print("\n1.Guyyaa | 2.Torbee | 3.Ji'aa")
                f = input("Filadhu: ")
                days = {'1':1, '2':7, '3':30}.get(f, 0)
                if days: uumi_excel(days, "Gabaasa")
            elif c == '3':
                # Asitti function barbaadi_galmee() waami
                pass
            elif c == '4': break
    else: print("[!] Login Dogoggora!")